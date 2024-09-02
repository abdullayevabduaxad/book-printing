import os

import PyPDF2
import docx
import openpyxl
from aiogram import Bot, Dispatcher, types, executor
from aiogram.contrib.fsm_storage.memory import MemoryStorage
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import State, StatesGroup
from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup, ParseMode
from config import token, admins, PAYMENT_TOKEN, Channel_token
from translations import translations
from sql import init_db, add_user, get_user, add_order
import re
import logging

logging.basicConfig()
# Bot va Dispatcher ni yaratamiz
TOKEN = token
bot = Bot(token=TOKEN)
storage = MemoryStorage()
dp = Dispatcher(bot, storage=storage)

# User orderlarini va tilni saqlash uchun dictionary
user_orders = {}
user_languages = {}


class Form(StatesGroup):
    fullname = State()
    phone = State()
    copy_count = State()


def get_translation(lang, key, **kwargs):
    translation = translations.get(lang, translations['en']).get(key, '')
    if not translation:
        return translations['en'].get(key, '')
    return translation.format(**kwargs)


@dp.message_handler(commands=['start'])
async def send_welcome(message: types.Message):
    user_id = message.from_user.id

    # Foydalanuvchi registratsiyadan o'tganligini tekshirish
    user_data = await get_user(user_id)  # Bu joyda get_user funksiyasi foydalanuvchini bazadan izlaydi
    if user_data:
        # Agar foydalanuvchi registratsiyadan o'tgan bo'lsa, menyuni ko'rsating
        btn = types.KeyboardButton('🛒 menu')
        aloqa = types.KeyboardButton('☎️ aloqa')
        about_me = types.KeyboardButton('✉️ about')
        royxat = types.KeyboardButton('ro`yxatdan o`tish')
        til = types.KeyboardButton('Tilni o`zgartirish')
        all_btn = types.ReplyKeyboardMarkup(resize_keyboard=True).add(btn, til).add(aloqa, about_me).add(royxat)
        welcome_msg = get_translation('uz', 'main')

        await message.reply(welcome_msg, reply_markup=all_btn)
    else:
        keyboard = [
            [
                InlineKeyboardButton("🇺🇿 Uzbek", callback_data='lang_uz'),
                InlineKeyboardButton("🇺🇸 English", callback_data='lang_en'),
                InlineKeyboardButton("🇷🇺 Русский", callback_data='lang_ru'),
            ],
        ]
        reply_markup = InlineKeyboardMarkup(inline_keyboard=keyboard)
        welcome_message = get_translation('en', 'select_language')  # Kalitni to'g'riladik
        if welcome_message:  # Tekshirish
            await message.reply(welcome_message, reply_markup=reply_markup)
        else:
            await message.reply("Error: Translation for 'select_language' not found.")


@dp.message_handler(text='Tilni o`zgartirish')
async def change_language(message: types.Message):
    keyboard = [
        [
            InlineKeyboardButton("🇺🇿 Uzbek", callback_data='lang_uz'),
            InlineKeyboardButton("🇺🇸 English", callback_data='lang_en'),
            InlineKeyboardButton("🇷🇺 Русский", callback_data='lang_ru'),
        ],
    ]
    reply_markup = InlineKeyboardMarkup(inline_keyboard=keyboard)
    welcome_message = get_translation('en', 'select_language')  # Kalitni to'g'riladik
    if welcome_message:  # Tekshirish
        await message.reply(welcome_message, reply_markup=reply_markup)
    else:
        await message.reply("Error: Translation for 'select_language' not found.")


@dp.message_handler(commands='register')
@dp.message_handler(text='ro`yxatdan o`tish')
async def cmd_register(message: types.Message):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    text = get_translation(lang, 'enter_fullname')
    if text:  # Matn bo'sh emasligini tekshirish
        await Form.fullname.set()
        await message.reply(text)
    else:
        await message.reply("Error: Translation for 'enter_fullname' not found.")


@dp.message_handler(state=Form.fullname)
async def process_fullname_step(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    text = get_translation(lang, 'enter_phone')
    if text:  # Matn bo'sh emasligini tekshirish
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True)
        button = types.KeyboardButton(get_translation(lang, 'share_contact_button'), request_contact=True)
        markup.add(button)

        async with state.proxy() as data:
            data['fullname'] = message.text

        await Form.next()
        await message.reply(text, reply_markup=markup)
    else:
        await message.reply("Error: Translation for 'enter_phone_number' not found.")


@dp.message_handler(state=Form.phone, content_types=types.ContentTypes.CONTACT)
async def process_phone_contact(message: types.Message, state: FSMContext):
    btn = types.KeyboardButton('🛒 menu')
    aloqa = types.KeyboardButton('☎️ aloqa')
    about_me = types.KeyboardButton('✉️ about')
    royxat = types.KeyboardButton('ro`yxatdan o`tish')
    til = types.KeyboardButton('Tilni o`zgartirish')
    all_btn = types.ReplyKeyboardMarkup(resize_keyboard=True).add(btn, til).add(aloqa, about_me).add(royxat)
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    text = get_translation(lang, 'registration_success')
    if text:  # Matn bo'sh emasligini tekshirish
        phone_number = message.contact.phone_number
        async with state.proxy() as data:
            data['phone'] = phone_number
            await add_user(user_id, data['fullname'], data['phone'])
        await message.reply(text, reply_markup=all_btn)
        await state.finish()
    else:
        await message.reply("Error: Translation for 'registration_successful' not found.")


@dp.message_handler(state=Form.phone)
async def process_phone_text(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    text = get_translation(lang, 'phone_request')
    btn = types.KeyboardButton('🛒 menu')
    aloqa = types.KeyboardButton('☎️ aloqa')
    about_me = types.KeyboardButton('about')
    royxat = types.KeyboardButton('ro`yxatdan o`tish')
    til = types.KeyboardButton('Tilni o`zgartirish')
    all_btn = types.ReplyKeyboardMarkup(resize_keyboard=True).add(btn, til).add(aloqa, about_me).add(royxat)
    if text:  # Matn bo'sh emasligini tekshirish
        markup = types.ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True)
        button = types.KeyboardButton(get_translation(lang, 'share_contact_button'), request_contact=True)
        markup.add(button)

        phone_number = message.text.strip()
        if re.match(r'^\+998\d{9}$', phone_number):
            async with state.proxy() as data:
                data['phone'] = phone_number
                await add_user(user_id, data['fullname'], data['phone'])
            await message.reply(get_translation(lang, 'registration_success'), reply_markup=all_btn)
            await state.finish()
        else:
            await message.reply(text, reply_markup=markup)
    else:
        await message.reply("Error: Translation for 'enter_phone_number' not found.")


def count_pdf_pages(file_path):
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            return len(reader.pages)
    except Exception as e:
        return f"Xato yuz berdi: {e}"


def count_docx_pages(file_path):
    try:
        doc = docx.Document(file_path)
        return len(doc.paragraphs) // 35  # Taxminiy har bir sahifada 35 paragraf
    except Exception as e:
        return f"Xato yuz berdi: {e}"


def count_xlsx_pages(file_path):
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        return (sheet.max_row // 50) + 1  # Taxminiy har bir sahifada 50 qator
    except Exception as e:
        return f"Xato yuz berdi: {e}"


def count_pages(file_path):
    if file_path.endswith('.pdf'):
        return count_pdf_pages(file_path)
    elif file_path.endswith('.docx'):
        return count_docx_pages(file_path)
    elif file_path.endswith('.xlsx'):
        return count_xlsx_pages(file_path)
    else:
        return "Ushbu format qo'llab-quvvatlanmaydi."


def get_format_keyboard(lang):
    keyboard = [
        [
            InlineKeyboardButton(get_translation(lang, 'a4'), callback_data='A4'),
            InlineKeyboardButton(get_translation(lang, 'a5'), callback_data='A5'),
        ],
    ]
    return InlineKeyboardMarkup(inline_keyboard=keyboard)


def get_color_keyboard(lang):
    keyboard = [
        [
            InlineKeyboardButton(get_translation(lang, 'color'), callback_data='color'),
            InlineKeyboardButton(get_translation(lang, 'black_white'), callback_data='black_white'),
        ],
    ]
    return InlineKeyboardMarkup(inline_keyboard=keyboard)


def get_method_keyboard(lang):
    keyboard = [
        [
            InlineKeyboardButton(get_translation(lang, 'thermal'), callback_data='thermal'),
            InlineKeyboardButton(get_translation(lang, 'normal'), callback_data='normal'),
        ],
    ]
    return InlineKeyboardMarkup(inline_keyboard=keyboard)


@dp.callback_query_handler(lambda c: c.data.startswith('lang_'))
async def process_language_choice(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    language_code = callback_query.data.split('_')[1]
    user_languages[user_id] = language_code

    user_data = await get_user(user_id)

    if user_data:
        # Foydalanuvchi registratsiyadan o'tgan bo'lsa, menyuni ko'rsatish
        btn = types.KeyboardButton('🛒 menu')
        aloqa = types.KeyboardButton('☎️ aloqa')
        about_me = types.KeyboardButton('✉️ about')
        royxat = types.KeyboardButton('ro`yxatdan o`tish')
        til = types.KeyboardButton('Tilni o`zgartirish')
        all_btn = types.ReplyKeyboardMarkup(resize_keyboard=True).add(btn, til).add(aloqa, about_me).add(royxat)

        await bot.send_message(chat_id=user_id, text=get_translation(language_code, 'main'), reply_markup=all_btn)
    else:
        # Foydalanuvchi registratsiyadan o'tmagan bo'lsa, registratsiya qilishni so'rash
        await bot.send_message(chat_id=user_id, text=get_translation(language_code, 'first_welcome'))

    await bot.answer_callback_query(callback_query.id)


@dp.message_handler(text='🛒 menu')
async def show_menu(message: types.Message):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa

    menu_message = get_translation(lang, 'menu_instruction')
    if menu_message:  # Tekshirish
        await message.reply(menu_message)
    else:
        await message.reply("Error: Translation for 'menu_instruction' not found.")


@dp.message_handler(text='✉️ about')
async def show_menu(message: types.Message):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')
    text = get_translation(lang, 'our')  # Foydalanuvchi tilini olish, default 'en' bo'lsa

    await message.answer(text)


@dp.message_handler(text='☎️ aloqa')
async def show_menu(message: types.Message):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa

    await message.answer('Bizning nomer!')


@dp.message_handler(text='tasdiqlash')
async def show_menu(message: types.Message):
    user_id = message.from_user.id

    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    user_order = user_orders.get(user_id)

    if not user_order:
        await bot.send_message(message.chat.id, "Buyurtma mavjud emas.")
        return

    total_cost = user_order['total_cost']
    copy_count = user_order['copy_count']
    total_payment = total_cost * copy_count  # Jami narx nusxa soniga ko'paytiriladi
    half_payment = int(total_payment * 0.5 * 100)  # 50% to'lov

    PRICE = types.LabeledPrice(label="Buyurtma uchun dastlabki to'lov.", amount=half_payment)
    await bot.send_invoice(chat_id=message.chat.id,
                           title="Buyurtma uchun dastlabki to'lov.",
                           description=f"Jami narxning 50%: {total_payment * 0.5} so'm.",
                           provider_token=PAYMENT_TOKEN[0],
                           currency="UZS",
                           prices=[PRICE],
                           start_parameter="usage-payment",
                           payload="usage-payment-payload")


@dp.pre_checkout_query_handler(lambda query: True)
async def checkout_process(query: types.PreCheckoutQuery):
    await bot.answer_pre_checkout_query(query.id, ok=True)


@dp.message_handler(content_types=types.ContentType.SUCCESSFUL_PAYMENT)
async def payment_success(message: types.Message):
    main = types.KeyboardButton('🛒 menu')
    aloqa = types.KeyboardButton('☎️ aloqa')
    about_me = types.KeyboardButton('✉️ about')
    royxat = types.KeyboardButton('ro`yxatdan o`tish')
    til = types.KeyboardButton('Tilni o`zgartirish')
    menu = types.ReplyKeyboardMarkup(resize_keyboard=True).add(main, til).add(aloqa, about_me).add(royxat)
    payment_info = message.successful_payment.to_python()
    for k, x in payment_info.items():
        print(f"{k}: {x}")
    await bot.send_message(message.chat.id, "To'lov qabul qilindi! Endi botdan foydalanishingiz mumkin.",
                           reply_markup=menu)
    await bot.send_message(admins[0], f"Foydalanuvchi {message.from_user.id} tomonidan to'lov qabul qilindi.")


@dp.message_handler(text='🔙 orqaga')
async def show_menu(message: types.Message):
    user_id = message.from_user.id
    lang = user_languages.get(user_id, 'en')  # Foydalanuvchi tilini olish, default 'en' bo'lsa
    bosh = types.KeyboardButton('🛒 menu')
    aloqa = types.KeyboardButton('☎️ aloqa')
    about_me = types.KeyboardButton('✉️ about')
    royxat = types.KeyboardButton('ro`yxatdan o`tish')
    til = types.KeyboardButton('Tilni o`zgartirish')
    all_btn = types.ReplyKeyboardMarkup(resize_keyboard=True).add(bosh, til).add(aloqa, about_me).add(royxat)
    await message.answer('bolimlardan birini tanla', reply_markup=all_btn)


@dp.message_handler(content_types=types.ContentType.DOCUMENT)
async def handle_document(message: types.Message):
    user_id = message.from_user.id
    if user_id not in user_languages:
        await bot.send_message(chat_id=user_id, text=get_translation('en', 'error_invalid_format'))
        return

    document = message.document
    file = await document.get_file()
    file_path = f"downloads/{user_id}_{document.file_name}"

    if not os.path.exists('downloads'):
        os.makedirs('downloads')

    await file.download(destination=file_path)

    user_orders[user_id] = {'file_path': file_path, 'file_name': document.file_name}

    await bot.send_message(chat_id=user_id, text=get_translation(user_languages[user_id], 'select_format'),
                           reply_markup=get_format_keyboard(user_languages[user_id]))


@dp.callback_query_handler(lambda c: c.data in ['A4', 'A5'])
async def process_format_choice(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id

    if user_id in user_orders:
        user_orders[user_id]['format'] = callback_query.data

        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'select_color'),
                               reply_markup=get_color_keyboard(user_languages[user_id]))
        await bot.answer_callback_query(callback_query.id)
    else:
        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'order_not_found'))
        await bot.answer_callback_query(callback_query.id)


@dp.callback_query_handler(lambda c: c.data in ['color', 'black_white'])
async def process_color_choice(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id

    if user_id in user_orders:
        user_orders[user_id]['color'] = callback_query.data

        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'select_method'),
                               reply_markup=get_method_keyboard(user_languages[user_id]))
        await bot.answer_callback_query(callback_query.id)
    else:
        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'order_not_found'))
        await bot.answer_callback_query(callback_query.id)


@dp.callback_query_handler(lambda c: c.data in ['thermal', 'normal'])
async def process_print_method_choice(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id

    if user_id in user_orders:
        user_orders[user_id]['method'] = callback_query.data

        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'enter_copy_count'))
        await Form.copy_count.set()
        await bot.answer_callback_query(callback_query.id)
    else:
        await bot.send_message(chat_id=callback_query.from_user.id,
                               text=get_translation(user_languages[user_id], 'order_not_found'))
        await bot.answer_callback_query(callback_query.id)


@dp.message_handler(state=Form.copy_count)
async def process_copy_count(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    copy_count = message.text.strip()

    if not copy_count.isdigit() or int(copy_count) <= 0:
        await message.reply(get_translation(user_languages[user_id], 'invalid_copy_count'))
        return

    user_orders[user_id]['copy_count'] = int(copy_count)

    file_path = user_orders[user_id]['file_path']
    format_choice = user_orders[user_id]['format']
    color_choice = user_orders[user_id]['color']
    method_choice = user_orders[user_id]['method']
    copy_count = user_orders[user_id]['copy_count']

    if format_choice == 'A4':
        price_per_page = 150
        binding_price = 7500 if method_choice == 'thermal' else 0
    elif format_choice == 'A5':
        price_per_page = 80
        binding_price = 7000 if method_choice == 'thermal' else 0

    price_multiplier = 1.5 if color_choice == 'color' else 1  # Rangli chop etish uchun narx koeffitsienti

    file_name = os.path.basename(file_path)
    file_size = count_pages(file_path)

    if isinstance(file_size, str):
        await bot.send_message(chat_id=user_id, text=f"Sahifalar sonini aniqlashda xato: {file_size}")
        return

    total_cost = (price_per_page * file_size + binding_price) * price_multiplier
    user_orders[user_id]['total_cost'] = total_cost
    total_payment = total_cost * copy_count
    response_text = (
        f"{get_translation(user_languages[user_id], 'order_received')}\n"
        f"{get_translation(user_languages[user_id], 'file')} {file_name}\n"
        f"{get_translation(user_languages[user_id], 'size')} {file_size} {get_translation(user_languages[user_id], 'pages')}\n"
        f"{get_translation(user_languages[user_id], 'format')} {format_choice}\n"
        f"{get_translation(user_languages[user_id], 'color')} {'Color' if color_choice == 'color' else 'Black and White'}\n"
        f"{get_translation(user_languages[user_id], 'method')} {'Thermal' if method_choice == 'thermal' else 'Normal'}\n"
        f"{get_translation(user_languages[user_id], 'copy_count')} {copy_count}\n"
        f"{get_translation(user_languages[user_id], 'price', price=total_cost)} x {copy_count} = {total_payment}\n"
        f"{get_translation(user_languages[user_id], 'payment_instruction')}"
    )

    user_data = await get_user(user_id)
    if user_data:
        fullname = user_data[0]
        phone = user_data[1]
    else:
        fullname = 'Noma\'lum'
        phone = 'Noma\'lum'

    adm = (
        f"Zakaz qabul qilindi:\n"
        f"Foydalanuvchi id: {user_id}\n"
        f"Fayl: {file_name}\n"
        f"Sahifalar soni: {file_size}\n"
        f"Format: {format_choice}\n"
        f"Rang: {'Rangli' if color_choice == 'color' else 'Qora va oq'}\n"
        f"Usul: {'Termal' if method_choice == 'thermal' else 'Normal'}\n"
        f"Nusxa soni: {copy_count}\n"
        f"Narx: {total_cost} so‘m\n"
        f"Foydalanuvchi ismi: {fullname}\n"
        f"Telefon raqam: {phone}\n"
    )

    btn = types.KeyboardButton('tasdiqlash')
    back_btn = types.KeyboardButton('🔙 orqaga')
    menu = types.ReplyKeyboardMarkup(resize_keyboard=True).add(btn, back_btn)

    await bot.send_message(chat_id=user_id, text=response_text, reply_markup=menu)
    for admin_id in admins:
        await bot.send_message(chat_id=admin_id, text=adm)
        await bot.send_document(chat_id=admin_id, document=open(file_path, 'rb'))
    chat = await bot.get_chat(Channel_token)
    channel_id = chat.id
    await bot.send_message(chat_id=channel_id, text=adm)
    await bot.send_document(chat_id=channel_id, document=open(file_path, 'rb'))
    os.remove(file_path)
    await state.finish()
    await add_order(user_id, file_path, file_name, format_choice, color_choice, method_choice, copy_count, total_cost)


async def on_startup(dp):
    await init_db()


if __name__ == '__main__':
    import asyncio

    loop = asyncio.get_event_loop()
    loop.run_until_complete(init_db())
    executor.start_polling(dp, skip_updates=True)
